#!/usr/bin/env python3
"""
Script to generate Property Brokerage documentation in DOCX format
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def add_heading_with_style(doc, text, level=1):
    """Add a styled heading to the document"""
    heading = doc.add_heading(text, level=level)
    return heading

def add_styled_paragraph(doc, text, bold=False, italic=False):
    """Add a styled paragraph to the document"""
    para = doc.add_paragraph()
    run = para.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    return para

def create_property_brokerage_documentation():
    """Create comprehensive Property Brokerage documentation"""
    
    # Create a new Document
    doc = Document()
    
    # Add title page
    title = doc.add_heading('Property Brokerage System', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Comprehensive Documentation')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_format = subtitle.runs[0]
    subtitle_format.font.size = Pt(16)
    subtitle_format.font.color.rgb = RGBColor(128, 128, 128)
    
    doc.add_paragraph()
    version_info = doc.add_paragraph('Version 1.0')
    version_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    date_info = doc.add_paragraph('January 2026')
    date_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Page break
    doc.add_page_break()
    
    # Table of Contents placeholder
    add_heading_with_style(doc, 'Table of Contents', level=1)
    doc.add_paragraph('1. Overview')
    doc.add_paragraph('2. System Features')
    doc.add_paragraph('3. System Architecture')
    doc.add_paragraph('4. User Roles and Permissions')
    doc.add_paragraph('5. Property Management')
    doc.add_paragraph('6. Client Management')
    doc.add_paragraph('7. Transaction Management')
    doc.add_paragraph('8. Reporting and Analytics')
    doc.add_paragraph('9. Technical Specifications')
    doc.add_paragraph('10. User Guide')
    doc.add_paragraph('11. Installation and Setup')
    doc.add_paragraph('12. API Documentation')
    doc.add_paragraph('13. Security and Compliance')
    doc.add_paragraph('14. Troubleshooting')
    doc.add_paragraph('15. Support and Contact Information')
    
    doc.add_page_break()
    
    # 1. Overview
    add_heading_with_style(doc, '1. Overview', level=1)
    doc.add_paragraph(
        'The Property Brokerage System is a comprehensive software solution designed to '
        'streamline and automate the operations of real estate brokerage firms. This system '
        'provides tools for managing properties, clients, agents, and transactions in a unified platform.'
    )
    
    add_heading_with_style(doc, '1.1 Purpose', level=2)
    doc.add_paragraph(
        'The primary purpose of this system is to:'
    )
    doc.add_paragraph('Centralize property listings and information', style='List Bullet')
    doc.add_paragraph('Manage client relationships and communications', style='List Bullet')
    doc.add_paragraph('Track real estate transactions from lead to closure', style='List Bullet')
    doc.add_paragraph('Provide analytics and reporting capabilities', style='List Bullet')
    doc.add_paragraph('Ensure compliance with real estate regulations', style='List Bullet')
    
    add_heading_with_style(doc, '1.2 Target Audience', level=2)
    doc.add_paragraph(
        'This system is designed for:'
    )
    doc.add_paragraph('Real estate brokers and agents', style='List Bullet')
    doc.add_paragraph('Property management companies', style='List Bullet')
    doc.add_paragraph('Real estate firms and agencies', style='List Bullet')
    doc.add_paragraph('Property owners and landlords', style='List Bullet')
    
    doc.add_page_break()
    
    # 2. System Features
    add_heading_with_style(doc, '2. System Features', level=1)
    
    add_heading_with_style(doc, '2.1 Property Listing Management', level=2)
    doc.add_paragraph('Create, edit, and manage property listings')
    doc.add_paragraph('Upload and organize property photos and documents')
    doc.add_paragraph('Set pricing, availability, and property details')
    doc.add_paragraph('Track property status (available, pending, sold)')
    
    add_heading_with_style(doc, '2.2 Client Relationship Management (CRM)', level=2)
    doc.add_paragraph('Maintain detailed client profiles and preferences')
    doc.add_paragraph('Track communication history and interactions')
    doc.add_paragraph('Manage leads and prospect conversion')
    doc.add_paragraph('Schedule appointments and follow-ups')
    
    add_heading_with_style(doc, '2.3 Transaction Management', level=2)
    doc.add_paragraph('Track deals from initial offer to closing')
    doc.add_paragraph('Manage contracts and legal documents')
    doc.add_paragraph('Calculate commissions and split payments')
    doc.add_paragraph('Generate transaction reports and summaries')
    
    add_heading_with_style(doc, '2.4 Search and Matching', level=2)
    doc.add_paragraph('Advanced property search with multiple filters')
    doc.add_paragraph('Automated client-property matching')
    doc.add_paragraph('Save searches and receive notifications')
    doc.add_paragraph('Map-based property visualization')
    
    add_heading_with_style(doc, '2.5 Reporting and Analytics', level=2)
    doc.add_paragraph('Sales performance dashboards')
    doc.add_paragraph('Agent productivity reports')
    doc.add_paragraph('Market trend analysis')
    doc.add_paragraph('Financial reports and forecasting')
    
    doc.add_page_break()
    
    # 3. System Architecture
    add_heading_with_style(doc, '3. System Architecture', level=1)
    
    doc.add_paragraph(
        'The Property Brokerage System follows a modern, scalable architecture designed '
        'for reliability, performance, and ease of maintenance.'
    )
    
    add_heading_with_style(doc, '3.1 Architecture Overview', level=2)
    doc.add_paragraph('Frontend: Web-based responsive interface')
    doc.add_paragraph('Backend: RESTful API service layer')
    doc.add_paragraph('Database: Relational database management system')
    doc.add_paragraph('Storage: Cloud-based file storage for media')
    doc.add_paragraph('Security: Authentication and authorization layer')
    
    add_heading_with_style(doc, '3.2 Technology Stack', level=2)
    doc.add_paragraph('Frontend Technologies:', style='List Bullet')
    doc.add_paragraph('  - HTML5, CSS3, JavaScript', style='List Bullet 2')
    doc.add_paragraph('  - Modern frontend framework (React/Vue/Angular)', style='List Bullet 2')
    doc.add_paragraph('  - Responsive design for mobile and desktop', style='List Bullet 2')
    
    doc.add_paragraph('Backend Technologies:', style='List Bullet')
    doc.add_paragraph('  - Server-side programming language', style='List Bullet 2')
    doc.add_paragraph('  - RESTful API architecture', style='List Bullet 2')
    doc.add_paragraph('  - Authentication and session management', style='List Bullet 2')
    
    doc.add_paragraph('Database:', style='List Bullet')
    doc.add_paragraph('  - Relational database (PostgreSQL/MySQL)', style='List Bullet 2')
    doc.add_paragraph('  - Data backup and recovery systems', style='List Bullet 2')
    
    doc.add_page_break()
    
    # 4. User Roles and Permissions
    add_heading_with_style(doc, '4. User Roles and Permissions', level=1)
    
    add_heading_with_style(doc, '4.1 Administrator', level=2)
    doc.add_paragraph('Full system access and configuration')
    doc.add_paragraph('User management and role assignment')
    doc.add_paragraph('System settings and customization')
    doc.add_paragraph('Access to all reports and analytics')
    
    add_heading_with_style(doc, '4.2 Broker/Manager', level=2)
    doc.add_paragraph('Manage agents and their assignments')
    doc.add_paragraph('Approve transactions and listings')
    doc.add_paragraph('Access team performance reports')
    doc.add_paragraph('Configure commission structures')
    
    add_heading_with_style(doc, '4.3 Agent', level=2)
    doc.add_paragraph('Create and manage property listings')
    doc.add_paragraph('Manage assigned clients and leads')
    doc.add_paragraph('Process transactions and deals')
    doc.add_paragraph('View personal performance metrics')
    
    add_heading_with_style(doc, '4.4 Client (Optional Portal Access)', level=2)
    doc.add_paragraph('Search available properties')
    doc.add_paragraph('Save favorite listings')
    doc.add_paragraph('Schedule property viewings')
    doc.add_paragraph('Track transaction status')
    
    doc.add_page_break()
    
    # 5. Property Management
    add_heading_with_style(doc, '5. Property Management', level=1)
    
    add_heading_with_style(doc, '5.1 Adding a Property', level=2)
    doc.add_paragraph('Navigate to the Property Management section')
    doc.add_paragraph('Click "Add New Property" button')
    doc.add_paragraph('Fill in required property details:')
    doc.add_paragraph('  - Property type (residential, commercial, land)', style='List Bullet 2')
    doc.add_paragraph('  - Address and location information', style='List Bullet 2')
    doc.add_paragraph('  - Size, bedrooms, bathrooms, amenities', style='List Bullet 2')
    doc.add_paragraph('  - Pricing and payment terms', style='List Bullet 2')
    doc.add_paragraph('Upload property photos and documents')
    doc.add_paragraph('Set property status and availability')
    doc.add_paragraph('Submit for approval (if required)')
    
    add_heading_with_style(doc, '5.2 Property Details', level=2)
    doc.add_paragraph('Basic Information: Address, size, type, year built')
    doc.add_paragraph('Financial Details: Price, taxes, HOA fees, financing options')
    doc.add_paragraph('Features: Bedrooms, bathrooms, parking, amenities')
    doc.add_paragraph('Media: Photos, videos, virtual tours, floor plans')
    doc.add_paragraph('Documents: Deeds, inspection reports, disclosures')
    
    add_heading_with_style(doc, '5.3 Property Status Management', level=2)
    doc.add_paragraph('Available: Property is on the market')
    doc.add_paragraph('Pending: Offer accepted, awaiting closing')
    doc.add_paragraph('Sold: Transaction completed')
    doc.add_paragraph('Off Market: Temporarily unavailable')
    doc.add_paragraph('Withdrawn: Removed from market')
    
    doc.add_page_break()
    
    # 6. Client Management
    add_heading_with_style(doc, '6. Client Management', level=1)
    
    add_heading_with_style(doc, '6.1 Client Profiles', level=2)
    doc.add_paragraph(
        'Client profiles store comprehensive information about buyers, sellers, and leads:'
    )
    doc.add_paragraph('Personal Information: Name, contact details, identification')
    doc.add_paragraph('Preferences: Property type, location, budget, features')
    doc.add_paragraph('Financial Information: Pre-approval status, financing options')
    doc.add_paragraph('Communication History: Calls, emails, meetings, notes')
    
    add_heading_with_style(doc, '6.2 Lead Management', level=2)
    doc.add_paragraph('Lead Capture: Web forms, phone calls, referrals')
    doc.add_paragraph('Lead Qualification: Assess buyer readiness and intent')
    doc.add_paragraph('Lead Assignment: Route leads to appropriate agents')
    doc.add_paragraph('Lead Nurturing: Automated follow-ups and communications')
    
    add_heading_with_style(doc, '6.3 Client Communication', level=2)
    doc.add_paragraph('Email integration and templates')
    doc.add_paragraph('SMS notifications and reminders')
    doc.add_paragraph('Appointment scheduling')
    doc.add_paragraph('Document sharing and signatures')
    
    doc.add_page_break()
    
    # 7. Transaction Management
    add_heading_with_style(doc, '7. Transaction Management', level=1)
    
    add_heading_with_style(doc, '7.1 Transaction Lifecycle', level=2)
    doc.add_paragraph('1. Lead Generation: Initial client contact')
    doc.add_paragraph('2. Property Showing: Schedule and conduct viewings')
    doc.add_paragraph('3. Offer Preparation: Draft and submit offers')
    doc.add_paragraph('4. Negotiation: Handle counter-offers and terms')
    doc.add_paragraph('5. Contract Execution: Sign purchase agreements')
    doc.add_paragraph('6. Due Diligence: Inspections, appraisals, financing')
    doc.add_paragraph('7. Closing: Final paperwork and fund transfer')
    doc.add_paragraph('8. Post-Closing: Follow-up and relationship maintenance')
    
    add_heading_with_style(doc, '7.2 Document Management', level=2)
    doc.add_paragraph('Contract templates and standard forms')
    doc.add_paragraph('Digital signature integration')
    doc.add_paragraph('Document version control')
    doc.add_paragraph('Secure document storage and access')
    
    add_heading_with_style(doc, '7.3 Commission Tracking', level=2)
    doc.add_paragraph('Commission rate configuration')
    doc.add_paragraph('Split calculations for multiple agents')
    doc.add_paragraph('Commission payment tracking')
    doc.add_paragraph('Year-to-date earnings reports')
    
    doc.add_page_break()
    
    # 8. Reporting and Analytics
    add_heading_with_style(doc, '8. Reporting and Analytics', level=1)
    
    add_heading_with_style(doc, '8.1 Available Reports', level=2)
    doc.add_paragraph('Sales Performance: Volume, value, conversion rates')
    doc.add_paragraph('Agent Activity: Listings, showings, closed deals')
    doc.add_paragraph('Property Analytics: Days on market, price trends')
    doc.add_paragraph('Client Reports: Lead sources, conversion funnels')
    doc.add_paragraph('Financial Reports: Revenue, commissions, expenses')
    
    add_heading_with_style(doc, '8.2 Dashboards', level=2)
    doc.add_paragraph('Executive Dashboard: High-level KPIs and metrics')
    doc.add_paragraph('Agent Dashboard: Personal performance and pipeline')
    doc.add_paragraph('Property Dashboard: Inventory and market insights')
    doc.add_paragraph('Financial Dashboard: Revenue and profitability')
    
    add_heading_with_style(doc, '8.3 Custom Reports', level=2)
    doc.add_paragraph('Report builder with drag-and-drop interface')
    doc.add_paragraph('Custom date ranges and filters')
    doc.add_paragraph('Export options (PDF, Excel, CSV)')
    doc.add_paragraph('Scheduled report delivery via email')
    
    doc.add_page_break()
    
    # 9. Technical Specifications
    add_heading_with_style(doc, '9. Technical Specifications', level=1)
    
    add_heading_with_style(doc, '9.1 System Requirements', level=2)
    doc.add_paragraph('Minimum Hardware:')
    doc.add_paragraph('  - Processor: 2 GHz dual-core or better', style='List Bullet 2')
    doc.add_paragraph('  - RAM: 4 GB minimum, 8 GB recommended', style='List Bullet 2')
    doc.add_paragraph('  - Storage: 20 GB available space', style='List Bullet 2')
    doc.add_paragraph('  - Network: Broadband internet connection', style='List Bullet 2')
    
    doc.add_paragraph('Software Requirements:')
    doc.add_paragraph('  - Modern web browser (Chrome, Firefox, Safari, Edge)', style='List Bullet 2')
    doc.add_paragraph('  - JavaScript enabled', style='List Bullet 2')
    doc.add_paragraph('  - PDF reader for document viewing', style='List Bullet 2')
    
    add_heading_with_style(doc, '9.2 Browser Compatibility', level=2)
    doc.add_paragraph('Google Chrome (latest 2 versions)')
    doc.add_paragraph('Mozilla Firefox (latest 2 versions)')
    doc.add_paragraph('Safari (latest 2 versions)')
    doc.add_paragraph('Microsoft Edge (latest 2 versions)')
    
    add_heading_with_style(doc, '9.3 Mobile Support', level=2)
    doc.add_paragraph('Responsive design for tablets and smartphones')
    doc.add_paragraph('iOS 12+ and Android 8+ support')
    doc.add_paragraph('Touch-optimized interface')
    doc.add_paragraph('Offline mode for essential features')
    
    doc.add_page_break()
    
    # 10. User Guide
    add_heading_with_style(doc, '10. User Guide', level=1)
    
    add_heading_with_style(doc, '10.1 Getting Started', level=2)
    doc.add_paragraph('1. Access the system through the provided URL')
    doc.add_paragraph('2. Log in with your username and password')
    doc.add_paragraph('3. Complete your profile information')
    doc.add_paragraph('4. Explore the dashboard and main navigation')
    doc.add_paragraph('5. Review the quick start tutorials')
    
    add_heading_with_style(doc, '10.2 Common Tasks', level=2)
    
    para = doc.add_paragraph()
    run = para.add_run('Adding a New Property:')
    run.bold = True
    doc.add_paragraph('1. Click "Properties" > "Add New"')
    doc.add_paragraph('2. Fill in all required fields')
    doc.add_paragraph('3. Upload photos and documents')
    doc.add_paragraph('4. Click "Save" or "Submit for Approval"')
    
    doc.add_paragraph()
    para = doc.add_paragraph()
    run = para.add_run('Creating a New Client:')
    run.bold = True
    doc.add_paragraph('1. Navigate to "Clients" > "Add Client"')
    doc.add_paragraph('2. Enter client information')
    doc.add_paragraph('3. Set preferences and budget')
    doc.add_paragraph('4. Assign to an agent')
    doc.add_paragraph('5. Save the client profile')
    
    doc.add_paragraph()
    para = doc.add_paragraph()
    run = para.add_run('Recording a Transaction:')
    run.bold = True
    doc.add_paragraph('1. Go to "Transactions" > "New Transaction"')
    doc.add_paragraph('2. Select property and client')
    doc.add_paragraph('3. Enter offer details and terms')
    doc.add_paragraph('4. Upload relevant documents')
    doc.add_paragraph('5. Track status through closing')
    
    doc.add_page_break()
    
    # 11. Installation and Setup
    add_heading_with_style(doc, '11. Installation and Setup', level=1)
    
    add_heading_with_style(doc, '11.1 Cloud-Based Deployment', level=2)
    doc.add_paragraph(
        'For cloud-based deployment, the system is typically hosted on cloud infrastructure '
        'and accessed via web browser. No local installation is required.'
    )
    doc.add_paragraph('1. Receive access credentials from administrator')
    doc.add_paragraph('2. Navigate to the system URL')
    doc.add_paragraph('3. Log in with provided credentials')
    doc.add_paragraph('4. Change password on first login')
    doc.add_paragraph('5. Complete initial setup wizard')
    
    add_heading_with_style(doc, '11.2 On-Premises Installation', level=2)
    doc.add_paragraph('For organizations requiring on-premises deployment:')
    doc.add_paragraph('1. Prepare server infrastructure')
    doc.add_paragraph('2. Install required software dependencies')
    doc.add_paragraph('3. Deploy application files')
    doc.add_paragraph('4. Configure database connection')
    doc.add_paragraph('5. Set up backup and security')
    doc.add_paragraph('6. Run initial configuration')
    doc.add_paragraph('7. Create administrator account')
    
    add_heading_with_style(doc, '11.3 Initial Configuration', level=2)
    doc.add_paragraph('Company Information: Name, logo, contact details')
    doc.add_paragraph('User Accounts: Create users and assign roles')
    doc.add_paragraph('Property Types: Define property categories')
    doc.add_paragraph('Commission Structure: Set default rates')
    doc.add_paragraph('Email Settings: Configure email integration')
    doc.add_paragraph('Document Templates: Upload standard forms')
    
    doc.add_page_break()
    
    # 12. API Documentation
    add_heading_with_style(doc, '12. API Documentation', level=1)
    
    doc.add_paragraph(
        'The Property Brokerage System provides a RESTful API for integration with '
        'third-party systems and custom applications.'
    )
    
    add_heading_with_style(doc, '12.1 API Authentication', level=2)
    doc.add_paragraph('Authentication Method: OAuth 2.0 / API Keys')
    doc.add_paragraph('API Endpoint: https://api.yourdomain.com/v1/')
    doc.add_paragraph('Required Headers:')
    doc.add_paragraph('  - Authorization: Bearer {token}', style='List Bullet 2')
    doc.add_paragraph('  - Content-Type: application/json', style='List Bullet 2')
    
    add_heading_with_style(doc, '12.2 Core API Endpoints', level=2)
    
    para = doc.add_paragraph()
    run = para.add_run('Properties API:')
    run.bold = True
    doc.add_paragraph('GET /properties - List all properties')
    doc.add_paragraph('GET /properties/{id} - Get property details')
    doc.add_paragraph('POST /properties - Create new property')
    doc.add_paragraph('PUT /properties/{id} - Update property')
    doc.add_paragraph('DELETE /properties/{id} - Delete property')
    
    doc.add_paragraph()
    para = doc.add_paragraph()
    run = para.add_run('Clients API:')
    run.bold = True
    doc.add_paragraph('GET /clients - List all clients')
    doc.add_paragraph('GET /clients/{id} - Get client details')
    doc.add_paragraph('POST /clients - Create new client')
    doc.add_paragraph('PUT /clients/{id} - Update client')
    
    doc.add_paragraph()
    para = doc.add_paragraph()
    run = para.add_run('Transactions API:')
    run.bold = True
    doc.add_paragraph('GET /transactions - List all transactions')
    doc.add_paragraph('GET /transactions/{id} - Get transaction details')
    doc.add_paragraph('POST /transactions - Create new transaction')
    doc.add_paragraph('PUT /transactions/{id} - Update transaction')
    
    doc.add_page_break()
    
    # 13. Security and Compliance
    add_heading_with_style(doc, '13. Security and Compliance', level=1)
    
    add_heading_with_style(doc, '13.1 Security Features', level=2)
    doc.add_paragraph('Encryption: Data encrypted in transit (SSL/TLS) and at rest')
    doc.add_paragraph('Authentication: Multi-factor authentication support')
    doc.add_paragraph('Authorization: Role-based access control (RBAC)')
    doc.add_paragraph('Audit Logs: Comprehensive activity logging')
    doc.add_paragraph('Session Management: Automatic timeout and session control')
    doc.add_paragraph('Password Policy: Enforced complexity and expiration')
    
    add_heading_with_style(doc, '13.2 Data Privacy', level=2)
    doc.add_paragraph('Compliance with data protection regulations (GDPR, CCPA)')
    doc.add_paragraph('Personal data encryption and anonymization')
    doc.add_paragraph('Right to access and delete personal data')
    doc.add_paragraph('Data retention policies and automated deletion')
    doc.add_paragraph('Privacy policy and terms of service')
    
    add_heading_with_style(doc, '13.3 Backup and Recovery', level=2)
    doc.add_paragraph('Automated daily backups')
    doc.add_paragraph('Point-in-time recovery capability')
    doc.add_paragraph('Off-site backup storage')
    doc.add_paragraph('Regular recovery testing')
    doc.add_paragraph('Disaster recovery plan')
    
    add_heading_with_style(doc, '13.4 Compliance', level=2)
    doc.add_paragraph('Real estate regulations compliance')
    doc.add_paragraph('Fair housing laws adherence')
    doc.add_paragraph('Financial transaction regulations')
    doc.add_paragraph('Industry standard security certifications')
    
    doc.add_page_break()
    
    # 14. Troubleshooting
    add_heading_with_style(doc, '14. Troubleshooting', level=1)
    
    add_heading_with_style(doc, '14.1 Common Issues', level=2)
    
    para = doc.add_paragraph()
    run = para.add_run('Login Problems:')
    run.bold = True
    doc.add_paragraph('Issue: Cannot log in to the system')
    doc.add_paragraph('Solution:')
    doc.add_paragraph('  - Verify username and password are correct', style='List Bullet 2')
    doc.add_paragraph('  - Check CAPS LOCK key', style='List Bullet 2')
    doc.add_paragraph('  - Use password reset if forgotten', style='List Bullet 2')
    doc.add_paragraph('  - Contact administrator if account is locked', style='List Bullet 2')
    
    doc.add_paragraph()
    para = doc.add_paragraph()
    run = para.add_run('File Upload Issues:')
    run.bold = True
    doc.add_paragraph('Issue: Unable to upload photos or documents')
    doc.add_paragraph('Solution:')
    doc.add_paragraph('  - Check file size (maximum 10 MB per file)', style='List Bullet 2')
    doc.add_paragraph('  - Verify file format is supported', style='List Bullet 2')
    doc.add_paragraph('  - Ensure stable internet connection', style='List Bullet 2')
    doc.add_paragraph('  - Try a different browser', style='List Bullet 2')
    
    doc.add_paragraph()
    para = doc.add_paragraph()
    run = para.add_run('Performance Issues:')
    run.bold = True
    doc.add_paragraph('Issue: System is slow or unresponsive')
    doc.add_paragraph('Solution:')
    doc.add_paragraph('  - Clear browser cache and cookies', style='List Bullet 2')
    doc.add_paragraph('  - Close unnecessary browser tabs', style='List Bullet 2')
    doc.add_paragraph('  - Check internet connection speed', style='List Bullet 2')
    doc.add_paragraph('  - Update to latest browser version', style='List Bullet 2')
    
    add_heading_with_style(doc, '14.2 Error Messages', level=2)
    doc.add_paragraph('Error 403 - Access Denied: Contact administrator for permissions')
    doc.add_paragraph('Error 404 - Not Found: Check URL or resource availability')
    doc.add_paragraph('Error 500 - Server Error: Contact technical support')
    doc.add_paragraph('Session Timeout: Log in again to continue')
    
    doc.add_page_break()
    
    # 15. Support and Contact Information
    add_heading_with_style(doc, '15. Support and Contact Information', level=1)
    
    add_heading_with_style(doc, '15.1 Technical Support', level=2)
    doc.add_paragraph('Email: support@propertybrokerage.com')
    doc.add_paragraph('Phone: 1-800-PROPERTY (1-800-776-7378)')
    doc.add_paragraph('Support Hours: Monday-Friday, 8:00 AM - 8:00 PM EST')
    doc.add_paragraph('Emergency Support: 24/7 for critical issues')
    
    add_heading_with_style(doc, '15.2 Additional Resources', level=2)
    doc.add_paragraph('Online Help Center: https://help.propertybrokerage.com')
    doc.add_paragraph('Video Tutorials: https://tutorials.propertybrokerage.com')
    doc.add_paragraph('User Community Forum: https://community.propertybrokerage.com')
    doc.add_paragraph('API Documentation: https://api.propertybrokerage.com/docs')
    
    add_heading_with_style(doc, '15.3 Training and Onboarding', level=2)
    doc.add_paragraph('Live training sessions available')
    doc.add_paragraph('Customized onboarding programs')
    doc.add_paragraph('Webinar schedule and recordings')
    doc.add_paragraph('Certification programs for advanced users')
    
    add_heading_with_style(doc, '15.4 Feedback and Suggestions', level=2)
    doc.add_paragraph('We value your feedback! Share your suggestions:')
    doc.add_paragraph('Email: feedback@propertybrokerage.com')
    doc.add_paragraph('Feature Request Portal: https://feedback.propertybrokerage.com')
    
    doc.add_page_break()
    
    # Appendix
    add_heading_with_style(doc, 'Appendix A: Glossary', level=1)
    
    terms = [
        ('Agent', 'A licensed real estate professional who represents buyers or sellers'),
        ('Broker', 'A licensed real estate professional who manages agents and transactions'),
        ('CRM', 'Customer Relationship Management - system for managing client interactions'),
        ('Closing', 'The final step in a real estate transaction where ownership transfers'),
        ('Commission', 'Percentage of sale price paid to agents/brokers for their services'),
        ('Lead', 'A potential client who has expressed interest in buying or selling'),
        ('Listing', 'A property that is available for sale or rent'),
        ('MLS', 'Multiple Listing Service - database of property listings'),
        ('Pre-approval', 'Lender\'s conditional agreement to provide financing'),
        ('Transaction', 'The complete process of buying or selling a property'),
    ]
    
    for term, definition in terms:
        para = doc.add_paragraph()
        run = para.add_run(f'{term}: ')
        run.bold = True
        para.add_run(definition)
    
    doc.add_page_break()
    
    # Document Information
    add_heading_with_style(doc, 'Document Information', level=1)
    doc.add_paragraph('Document Title: Property Brokerage System Documentation')
    doc.add_paragraph('Version: 1.0')
    doc.add_paragraph('Date: January 2026')
    doc.add_paragraph('Status: Final')
    doc.add_paragraph('Classification: Internal Use')
    
    doc.add_paragraph()
    add_heading_with_style(doc, 'Revision History', level=2)
    
    # Create a table for revision history
    table = doc.add_table(rows=2, cols=4)
    table.style = 'Light Grid Accent 1'
    
    # Header row
    header_cells = table.rows[0].cells
    header_cells[0].text = 'Version'
    header_cells[1].text = 'Date'
    header_cells[2].text = 'Author'
    header_cells[3].text = 'Description'
    
    # Data row
    data_cells = table.rows[1].cells
    data_cells[0].text = '1.0'
    data_cells[1].text = 'January 2026'
    data_cells[2].text = 'System Team'
    data_cells[3].text = 'Initial documentation release'
    
    # Save the document
    output_path = '/home/runner/work/Property-Brokerage-/Property-Brokerage-/Property_Brokerage_Documentation.docx'
    doc.save(output_path)
    print(f'Documentation successfully created: {output_path}')
    
    return output_path

if __name__ == '__main__':
    create_property_brokerage_documentation()
